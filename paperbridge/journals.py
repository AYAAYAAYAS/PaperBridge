"""DOI journal identification, authorized ranking providers and local annual cache."""
from __future__ import annotations

import base64
import csv
import io
import json
import re
import threading
import unicodedata
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import httpx
import pymupdf

from paperbridge.core import _dpapi, atomic_json
from paperbridge.network import ssl_context

CAS_API = "https://webapi.fenqubiao.com/api/user/get"
JCR_API = "https://api.clarivate.com/apis/wos-journals/v1"
HEADERS = ["期刊名称", "ISSN", "体系", "年份", "分区", "学科", "级别", "来源"]


def timestamp():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def journal_key(value):
    return re.sub(r"[^\w]", "", unicodedata.normalize("NFKC", str(value)).casefold())


def issns(value):
    if isinstance(value, list):
        value = " ".join(str(x) for x in value)
    return sorted(set(x.upper() for x in re.findall(r"\b\d{4}-\d{3}[\dXx]\b", str(value))))


def same_journal(journal, identifiers, record):
    left, right = set(issns(identifiers)), set(issns(record.get("issns", [])))
    if left and right:
        return bool(left & right)
    return bool(journal_key(journal)) and journal_key(journal) == journal_key(record.get("journal", ""))


def pdf_doi(path):
    try:
        with pymupdf.open(path) as doc:
            for index in range(min(2, len(doc))):
                page = doc[index]
                # Publisher DOI hyperlinks precede incidental DOI text on page one.
                text = " ".join(x.get("uri", "") for x in page.get_links()) + "\n" + page.get_text()
                match = re.search(r"10\.\d{4,9}/[^\s<>\[\]\"']+", text, re.I)
                if match:
                    return match.group().rstrip(".,;:)}")
    except (OSError, RuntimeError, ValueError):
        pass
    return ""


def parse_cas(payload, journal, identifiers, year):
    if not isinstance(payload, dict) or int(payload.get("Year", 0)) != year:
        raise ValueError("中科院接口未返回指定年份的数据。")
    identity = {"journal": payload.get("Title", ""), "issns": issns(payload.get("ISSN", ""))}
    if not same_journal(journal, identifiers, identity):
        raise ValueError("中科院接口的期刊标识不一致，结果未采用。")
    result = []
    # The CAS provider's JCR field is CAS small-category ranking, NOT Clarivate Q1-Q4.
    for field, scope in (("ZKY", "大类"), ("JCR", "小类")):
        for category in payload.get(field) or []:
            zone = str(category.get("Section", ""))
            if zone in ("1", "2", "3", "4"):
                result.append(dict(identity, scheme="CAS", year=year, scope=scope,
                                   category=category.get("NameCN") or category.get("Name", ""),
                                   zone=zone + "区", source="中科院期刊分区表授权接口",
                                   source_url=CAS_API, fetched_at=timestamp()))
    return result


def parse_jcr(payload, identity, year):
    if int(payload.get("year", 0)) != year:
        raise ValueError("JCR 接口未返回指定年份的数据。")
    result = []
    for rank in (payload.get("ranks") or {}).get("jif") or []:
        zone = str(rank.get("quartile", "")).upper()
        if zone in ("Q1", "Q2", "Q3", "Q4"):
            result.append(dict(identity, scheme="JCR", year=year, scope="JIF",
                               category=rank.get("category", ""), zone=zone,
                               source="Clarivate JCR 授权接口", source_url=JCR_API,
                               fetched_at=timestamp()))
    return result


def ranking_summary(paper):
    records = paper.get("rankings") or []
    lines = []
    for scheme, label in (("CAS", "中科院"), ("JCR", "JCR")):
        selected = [r for r in records if r["scheme"] == scheme]
        if selected:
            parts = [f"{r['scope']} {r['category']} {r['zone']}".strip() for r in selected]
            lines.append(f"{label} {selected[0]['year']}：" + "；".join(parts))
        else:
            year = paper.get("ranking_year", "")
            lines.append(f"{label} {year}：未获取".replace("  ", " "))
    if paper.get("partition"):
        lines.append("原手工备注：" + paper["partition"])
    return "\n".join(lines)


def ranking_brief(paper):
    parts = []
    for scheme, label in (("CAS", "中科院"), ("JCR", "JCR")):
        records = [r for r in paper.get("rankings", []) if r["scheme"] == scheme]
        if records:
            zones = list(dict.fromkeys((r["scope"] + r["zone"] if scheme == "CAS" else r["zone"]) for r in records))
            parts.append(f"{label} {records[0]['year']}：" + " / ".join(zones))
        else:
            parts.append(f"{label}：未获取")
    if paper.get("partition"):
        parts.append("备注：" + paper["partition"])
    return "\n".join(parts)


class LookupCancelled(Exception):
    pass


class JournalStore:
    def __init__(self, root):
        self.root = Path(root)
        self.cache_file = self.root / "journal_cache.json"
        self.config_file = self.root / "journal_sources.json"
        self.lock = threading.RLock()
        self.cache = self._read(self.cache_file, {"dois": {}, "rankings": []})
        if not isinstance(self.cache.get("dois"), dict) or not isinstance(self.cache.get("rankings"), list):
            self.cache = {"dois": {}, "rankings": []}
        self.config = dict(year=datetime.now().year - 1, auto=True, remember=False,
                           cas_user="", cas_password="", jcr_key="")
        stored = self._read(self.config_file, {})
        self.config.update({k: stored[k] for k in ("year", "auto", "remember") if k in stored})
        if stored.get("protected"):
            try:
                self.config.update(json.loads(_dpapi(base64.b64decode(stored["protected"]), True)))
            except (OSError, ValueError, UnicodeError):
                pass

    @staticmethod
    def _read(path, default):
        try:
            result = json.loads(path.read_text(encoding="utf-8"))
            return result if isinstance(result, dict) else default
        except (OSError, ValueError):
            return default

    def configure(self, values):
        stored = {k: values[k] for k in ("year", "auto", "remember")}
        if values["remember"]:
            secret = {k: values[k] for k in ("cas_user", "cas_password", "jcr_key")}
            stored["protected"] = base64.b64encode(_dpapi(json.dumps(secret).encode())).decode()
        atomic_json(self.config_file, stored)
        self.config.update(values)

    def cached(self, journal, identifiers, year):
        with self.lock:
            return [dict(r) for r in self.cache["rankings"] if r["year"] == year
                    and same_journal(journal, identifiers, r)]

    def store_records(self, records):
        with self.lock:
            current = list(self.cache["rankings"])
            for record in records:
                current = [r for r in current if not (
                    r["scheme"] == record["scheme"] and r["year"] == record["year"]
                    and r["scope"] == record["scope"] and r["category"] == record["category"]
                    and same_journal(record["journal"], record["issns"], r))]
                current.append(record)
            atomic_json(self.cache_file, dict(self.cache, rankings=current))
            self.cache["rankings"] = current

    def lookup(self, paper, force=False, cancelled=lambda: False):
        config = dict(self.config)
        year = int(config["year"])
        notes = []

        def request(url, params=None, headers=None):
            if cancelled():
                raise LookupCancelled()
            try:
                with httpx.Client(trust_env=False, verify=ssl_context(), timeout=12) as client:
                    response = client.get(url, params=params, headers=headers or {})
                response.raise_for_status()
                return response.json()
            except httpx.HTTPStatusError as exc:
                raise ValueError(f"查询服务返回 HTTP {exc.response.status_code}，请检查授权和年份。") from None
            except (httpx.RequestError, ValueError):
                # Never include a request URL: CAS places its credentials in query parameters.
                raise ValueError("查询服务连接失败或响应异常，请检查网络及数据源授权。") from None

        metadata = {k: paper.get(k, "") for k in ("journal", "doi")}
        metadata["issns"] = issns(paper.get("issns", []))
        path = paper.get("managed")
        metadata["doi"] = metadata["doi"] or pdf_doi(path if path and Path(path).is_file() else paper["source"])
        doi = metadata["doi"]
        if doi:
            with self.lock:
                identified = self.cache["dois"].get(doi.lower())
            if not identified or force:
                try:
                    response = request("https://api.crossref.org/works/" + quote(doi, safe="/"))
                    message = response["message"]
                    titles = message.get("container-title") or []
                    if titles and message.get("type") == "journal-article":
                        identified = {"journal": titles[0], "issns": issns(message.get("ISSN", [])),
                                      "doi": doi, "metadata_source": "Crossref DOI", "identified_at": timestamp()}
                        with self.lock:
                            self.cache["dois"][doi.lower()] = identified
                            atomic_json(self.cache_file, self.cache)
                    else:
                        notes.append("该 DOI 未对应期刊论文，未推断期刊。")
                except (ValueError, KeyError, TypeError):
                    notes.append("DOI 期刊识别失败；可手动填写期刊 / ISSN 后查询。")
            if identified:
                if metadata["journal"] and journal_key(metadata["journal"]) != journal_key(identified["journal"]):
                    notes.append("填写期刊与 DOI 期刊不同，保留手工期刊；请核对 ISSN。")
                else:
                    metadata.update(identified)
        if not metadata["journal"] and not metadata["issns"]:
            return {"metadata": metadata, "rankings": [], "year": year,
                    "status": "未识别到期刊 DOI；请先填写期刊全名或 ISSN。"}
        journal, identifiers = metadata["journal"], metadata["issns"]
        for scheme in ("CAS", "JCR"):
            if not force and any(r["scheme"] == scheme for r in self.cached(journal, identifiers, year)):
                continue
            try:
                records = []
                keyword = next(iter(identifiers), journal)
                if scheme == "CAS" and config["cas_user"] and config["cas_password"]:
                    response = request(CAS_API, {"year": year, "keyword": keyword,
                                                "user": config["cas_user"], "password": config["cas_password"]})
                    records = parse_cas(response, journal, identifiers, year)
                elif scheme == "JCR" and config["jcr_key"]:
                    headers = {"X-ApiKey": config["jcr_key"]}
                    response = request(JCR_API + "/journals", {"q": keyword, "limit": 5}, headers)
                    for hit in (response.get("hits") or [])[:3]:
                        detail = request(JCR_API + "/journals/" + quote(hit["id"], safe=""), headers=headers)
                        identity = {"journal": detail.get("name", ""),
                                    "issns": issns([detail.get("issn", ""), detail.get("eIssn", "")])}
                        if same_journal(journal, identifiers, identity):
                            report = request(JCR_API + "/journals/" + quote(hit["id"], safe="")
                                             + f"/reports/year/{year}", headers=headers)
                            records = parse_jcr(report, identity, year)
                            break
                else:
                    notes.append(f"{'中科院' if scheme == 'CAS' else 'JCR'}：未配置授权，可导入分区表。")
                    continue
                if records:
                    self.store_records(records)
                else:
                    notes.append(f"{scheme}：该期刊在 {year} 年未返回可用分区。")
            except (ValueError, KeyError, TypeError) as exc:
                notes.append(f"{scheme}：{exc}" if isinstance(exc, ValueError) else f"{scheme}：接口响应格式异常。")
        records = self.cached(journal, identifiers, year)
        if not metadata["journal"] and records:
            metadata["journal"] = records[0]["journal"]
        return {"metadata": metadata, "rankings": records,
                "year": year, "status": "；".join(notes) or "已从本地缓存读取 / 查询完成。"}

    def import_table(self, path):
        path = Path(path)
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = list(book.worksheets[0].values)
            finally:
                book.close()
        else:
            raw = path.read_bytes()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("gb18030")
            rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            raise ValueError("分区表为空。")
        aliases = {"journal": ["期刊名称", "期刊", "journal", "title"],
                   "issn": ["ISSN", "issn", "eISSN"], "scheme": ["体系", "scheme"],
                   "year": ["年份", "year"], "zone": ["分区", "zone", "quartile"],
                   "category": ["学科", "category"], "scope": ["级别", "scope"],
                   "source": ["来源", "source"]}
        headers = [str(x or "").strip() for x in rows[0]]
        columns = {key: next((i for i, h in enumerate(headers) if h.casefold() in [a.casefold() for a in names]), None)
                   for key, names in aliases.items()}
        if any(columns[k] is None for k in ("scheme", "year", "zone")):
            raise ValueError("缺少表头：体系、年份、分区。请先导出模板并按格式填写。")
        result = []
        for number, row in enumerate(rows[1:], 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            def cell(key):
                index = columns[key]
                return str(row[index] if index is not None and index < len(row) and row[index] is not None else "").strip()
            scheme = {"中科院": "CAS", "cas": "CAS", "jcr": "JCR"}.get(cell("scheme").lower())
            try:
                value = float(cell("year"))
                year = int(value) if value.is_integer() else 0
            except ValueError:
                year = 0
            zone = cell("zone").upper().replace(" ", "")
            if scheme == "CAS":
                zone = zone.replace("区", "")
                zone = zone + "区" if zone in ("1", "2", "3", "4") else ""
            elif scheme == "JCR":
                zone = zone if zone in ("Q1", "Q2", "Q3", "Q4") else ""
            scope = cell("scope") or ("大类" if scheme == "CAS" else "JIF")
            if scheme == "JCR":
                scope = scope.upper()
            if (not scheme or not 2000 <= year <= 2100 or not zone
                    or not (cell("journal") or issns(cell("issn")))
                    or (scheme == "CAS" and scope not in ("大类", "小类"))
                    or (scheme == "JCR" and scope.upper() != "JIF")):
                raise ValueError(f"第 {number} 行的体系、年份、分区、级别或期刊标识无效，未导入任何行。")
            record = {"journal": cell("journal"), "issns": issns(cell("issn")),
                           "scheme": scheme, "year": year, "zone": zone, "scope": scope,
                           "category": cell("category"), "source": "用户导入：" + (cell("source") or path.name),
                           "source_url": "", "fetched_at": timestamp()}
            if any(r["scheme"] == scheme and r["year"] == year and r["scope"] == scope
                   and r["category"] == record["category"] and r["zone"] != zone
                   and same_journal(record["journal"], record["issns"], r) for r in result):
                raise ValueError(f"第 {number} 行与前面的同刊、同年、同学科记录分区冲突，未导入任何行。")
            result.append(record)
        if not result:
            raise ValueError("表格只有表头，尚无可导入的分区数据。")
        self.store_records(result)
        return len(result)

    @staticmethod
    def export_template(path):
        path = Path(path)
        if path.suffix.lower() == ".xlsx":
            from openpyxl import Workbook
            book = Workbook()
            sheet = book.active
            sheet.title = "分区数据"
            sheet.append(HEADERS)
            sheet.freeze_panes = "A2"
            for column in "ABCDEFGH":
                sheet.column_dimensions[column].width = 24
            book.save(path)
        else:
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                csv.writer(handle).writerow(HEADERS)
